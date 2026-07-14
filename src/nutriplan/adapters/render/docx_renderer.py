"""Renderer DOCX editable (python-docx) detrás del mismo puerto Renderer."""

from io import BytesIO
from uuid import UUID

from nutriplan.adapters.render.view import (
    DAY_LABELS,
    PHASE_SUBTITLES,
    SLOT_LABELS,
    anotaciones,
    free_meal_cell,
    build_grid,
    day_totals_row,
    macro_line,
    plan_phases_in,
    slots_in,
)
from nutriplan.domain.errors import RenderError
from nutriplan.domain.models import Branding, FoodItem, MacroTargets, PlanCycle


class DocxRenderer:
    """Implementa el puerto Renderer para fmt='docx'."""

    async def render(
        self,
        plan: PlanCycle,
        branding: Branding,
        foods: dict[UUID, FoodItem],
        fmt: str = "docx",
        client_name: str | None = None,
        daily_targets: MacroTargets | None = None,
    ) -> bytes:
        if fmt != "docx":
            raise RenderError(f"DocxRenderer solo produce docx, no {fmt}")
        try:
            return self._build(plan, branding, foods, daily_targets=daily_targets)
        except RenderError:
            raise
        except Exception as exc:
            raise RenderError(f"python-docx falló: {exc}") from exc

    def _build(
        self,
        plan: PlanCycle,
        branding: Branding,
        foods: dict[UUID, FoodItem],
        *,
        daily_targets: MacroTargets | None = None,
    ) -> bytes:
        from docx import Document
        from docx.enum.section import WD_ORIENT

        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

        doc.add_heading(f"{branding.tenant_name} — Plan nutricional", level=1)
        daily = daily_targets
        if daily is None and plan.days:
            n = len(plan.days)
            daily = MacroTargets(
                kcal=sum(d.totals.kcal for d in plan.days) / n,
                protein_g=sum(d.totals.protein_g for d in plan.days) / n,
                carb_g=sum(d.totals.carb_g for d in plan.days) / n,
                fat_g=sum(d.totals.fat_g for d in plan.days) / n,
            )
        if daily is not None:
            doc.add_paragraph(f"Objetivo diario · {macro_line(daily)}")

        phases = plan_phases_in(plan)
        multi = len(phases) > 1 or plan.duration_days >= 30
        slots = slots_in(plan)  # las comidas que este cliente come, no siempre 5
        for phase in phases:
            if multi:
                doc.add_heading(PHASE_SUBTITLES.get(phase, phase.value), level=2)
            grid = build_grid(plan, foods, phase=phase)
            totals = day_totals_row(plan, phase)
            table = doc.add_table(rows=2 + len(slots), cols=8)
            table.style = "Table Grid"

            header = table.rows[0].cells
            header[0].text = ""
            for i, day in enumerate(DAY_LABELS, start=1):
                header[i].text = day

            for r, (slot, cells) in enumerate(grid.items(), start=1):
                row = table.rows[r].cells
                row[0].text = SLOT_LABELS[slot]
                for c, cell in enumerate(cells, start=1):
                    lines = [p.text for p in cell.portions] + cell.extras
                    if cell.macro_line:
                        lines.append(cell.macro_line)
                    row[c].text = "\n".join(lines)

            total_row = table.rows[1 + len(slots)].cells
            total_row[0].text = "Totales día"
            for c, dt in enumerate(totals, start=1):
                total_row[c].text = (
                    f"{dt.kcal} kcal\n"
                    f"P {dt.protein_g} g · C {dt.carb_g} g · G {dt.fat_g} g"
                )

        doc.add_heading("Anotaciones Importantes", level=2)
        for nota in anotaciones(slots, free_meal_cell(plan)):
            doc.add_paragraph(nota, style="List Number")

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
