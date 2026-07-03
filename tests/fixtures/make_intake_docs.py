"""Genera los .docx sintéticos de intake (correr una vez; artefactos en git).

    uv run python tests/fixtures/make_intake_docs.py
"""

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent / "intakes"


def _doc(title: str, lines: list[str]) -> Document:
    doc = Document()
    doc.add_heading(title, level=1)
    for line in lines:
        doc.add_paragraph(line)
    return doc


def main() -> None:
    OUT.mkdir(exist_ok=True)

    _doc(
        "Cuestionario inicial — Ana Pérez",
        [
            "Nombre completo: Ana Pérez",
            "Sexo: femenino",
            "Edad: 28 años",
            "Ciudad: Medellín",
            "Estatura: 1.65 m",
            "Peso actual: 62 kg",
            "Objetivo: bajar grasa y tonificar",
            "Entrenamiento: pesas 4 veces por semana",
            "Proteínas que le gustan: pollo, huevos, atún",
            "Carbohidratos: arroz, arepa, papa",
            "Grasas: aguacate, maní",
            "Frutas: banano, fresa, mango",
            "Verduras: brócoli, espinaca, tomate",
            "Restricciones: sin mariscos",
            "¿Usa batido de proteína?: no",
        ],
    ).save(OUT / "intake_limpio.docx")

    _doc(
        "Cuestionario inicial — Carlos Ruiz",
        [
            "Nombre completo: Carlos Ruiz",
            "Sexo: masculino",
            "Edad: 35 años",
            "Ciudad: Bogotá",
            "Estatura: 178 cm",
            "Peso: 84 kg aproximado, no me he pesado esta semana, lo confirmo",
            "Objetivo: ganar masa muscular",
            "Entrenamiento: gimnasio 5 días",
            "Proteínas: carne de res, pollo, huevos",
            "Carbohidratos: pasta, arroz, pan",
            "Grasas: aceite de oliva, almendras",
            "Frutas: manzana, uvas",
            "Verduras: lechuga, zanahoria",
            "Restricciones: ninguna",
            "¿Usa batido de proteína?: sí, whey después de entrenar",
        ],
    ).save(OUT / "intake_peso_aproximado.docx")

    _doc(
        "Cuestionario inicial — Laura Gómez",
        [
            "Nombre completo: Laura Gómez",
            "Sexo: femenino",
            "Edad: 42 años",
            "Peso: 70 kg",
            # sin estatura — intake incompleto a propósito
            "Objetivo: mantenerme y comer mejor",
            "Proteínas: tilapia, pechuga de pavo",
            "Carbohidratos: quinoa, batata",
            "Grasas: nueces",
            "Frutas: papaya, kiwi",
            "Verduras: coliflor, pepino",
            "Restricciones: sin gluten, sin lácteos",
        ],
    ).save(OUT / "intake_incompleto.docx")

    print(f"Fixtures generados en {OUT}")


if __name__ == "__main__":
    main()
